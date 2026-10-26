import os
import re
import io
import sqlite3
import glob
import hashlib
import mimetypes
from openpyxl.styles import Alignment, Font, PatternFill
from urllib.parse import urlparse
from urllib.request import Request, urlopen
import pandas as pd
from datetime import datetime
from flask import Flask, request, jsonify, send_from_directory, send_file
from flask_cors import CORS
from werkzeug.utils import secure_filename
from concurrent.futures import ThreadPoolExecutor

app = Flask(__name__)
CORS(app)

DB_NAME = "database.db"
UPLOAD_FOLDER = 'uploads'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

IMAGE_CACHE = os.path.abspath(os.path.join(UPLOAD_FOLDER, 'cache'))
os.makedirs(IMAGE_CACHE, exist_ok=True)

def transform_drive_url(url):
    """Google Drive view link ko direct image render link me convert karta hai"""
    if not isinstance(url, str) or not url.strip():
        return ""
    match = re.search(r'(?:id=|\/d\/)([a-zA-Z0-9_-]+)', url)
    if match:
        return f"https://lh3.googleusercontent.com/d/{match.group(1)}"
    return url.strip()

def get_col_val(row, possible_names, default="--"):
    """Excel headers me case-insensitive matching karta hai"""
    def normalize_header(value):
        return re.sub(r'[^a-z0-9]', '', str(value).strip().lower())

    for col in row.index:
        clean_col = normalize_header(col)
        for name in possible_names:
            if clean_col == normalize_header(name):
                val = row[col]
                if pd.notna(val) and str(val).strip() != "":
                    return str(val).strip()
    return default

def format_reg_no(value):
    """Format numeric registration numbers as three digits."""
    value = str(value).strip()
    if re.fullmatch(r'\d+(?:\.0+)?', value):
        return str(int(float(value))).zfill(3)
    return value

def migrate_registration_numbers():
    conn = sqlite3.connect(DB_NAME)
    rows = conn.execute("SELECT id, reg_no FROM players ORDER BY id").fetchall()
    for player_id, reg_no in rows:
        formatted = format_reg_no(reg_no)
        if formatted != reg_no:
            conn.execute("UPDATE players SET reg_no=? WHERE id=?", (formatted, player_id))
    conn.commit()
    conn.close()

def next_registration_number(used_numbers):
    number = 1
    while format_reg_no(number) in used_numbers:
        number += 1
    registration_number = format_reg_no(number)
    used_numbers.add(registration_number)
    return registration_number

def init_db():
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    
    # Players Table (added block column)
    c.execute('''CREATE TABLE IF NOT EXISTS players (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        reg_no TEXT UNIQUE,
        name TEXT,
        category TEXT,
        batting_style TEXT,
        bowling_style TEXT,
        district TEXT,
        block TEXT,
        photo_url TEXT,
        base_price REAL DEFAULT 400,
        status TEXT DEFAULT 'PENDING',
        sold_to TEXT DEFAULT NULL,
        sold_price REAL DEFAULT 0,
        sold_at DATETIME DEFAULT NULL,
        round_no INTEGER DEFAULT 1
    )''')
    
    # Safe migration for databases created before sold_at was added.
    player_columns = {row[1] for row in c.execute("PRAGMA table_info(players)").fetchall()}
    if 'sold_at' not in player_columns:
        c.execute("ALTER TABLE players ADD COLUMN sold_at DATETIME DEFAULT NULL")

    c.execute('''CREATE TABLE IF NOT EXISTS teams (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT UNIQUE,
        budget REAL DEFAULT 50000,
        purse_left REAL DEFAULT 50000,
        squad_count INTEGER DEFAULT 0
    )''')

    c.execute('''CREATE TABLE IF NOT EXISTS auction_settings (
        id INTEGER PRIMARY KEY DEFAULT 1,
        bid_step_1 REAL DEFAULT 100,
        bid_step_2 REAL DEFAULT 500,
        bid_step_3 REAL DEFAULT 1000,
        current_round INTEGER DEFAULT 1
    )''')

    c.execute('''CREATE TABLE IF NOT EXISTS bids_history (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        player_id INTEGER,
        team_name TEXT,
        bid_amount REAL,
        timestamp DATETIME DEFAULT CURRENT_TIMESTAMP
    )''')

    c.execute("INSERT OR IGNORE INTO auction_settings (id, bid_step_1, bid_step_2, bid_step_3, current_round) VALUES (1, 100, 500, 1000, 1)")
    conn.commit()
    conn.close()
    migrate_registration_numbers()

init_db()

@app.route('/uploads/<filename>')
def uploaded_file(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)

def drive_source_url(photo_url):
    """Player photo link -> the exact Google image URL the frontend asks for."""
    match = re.search(r'/d/([a-zA-Z0-9_-]+)', photo_url or '')
    return f"https://lh3.googleusercontent.com/d/{match.group(1)}=w800" if match else None

def find_cached_image(source_url):
    key = hashlib.sha1(source_url.encode('utf-8')).hexdigest()
    found = glob.glob(os.path.join(IMAGE_CACHE, key + '.*'))
    return found[0] if found else None

def download_image(source_url):
    """Downloads one image from Google and saves it on disk. Returns the saved file path."""
    upstream_request = Request(source_url, headers={'User-Agent': 'LiveAuctions/1.0'})
    with urlopen(upstream_request, timeout=10) as upstream:
        image_data = upstream.read()
        content_type = upstream.headers.get_content_type()
    if not content_type.startswith('image/'):
        raise ValueError('Google did not return an image (is the Drive file shared as Anyone with the link?)')
    key = hashlib.sha1(source_url.encode('utf-8')).hexdigest()
    path = os.path.join(IMAGE_CACHE, key + (mimetypes.guess_extension(content_type) or '.jpg'))
    with open(path, 'wb') as image_file:
        image_file.write(image_data)
    return path

@app.route('/api/image-proxy')
def image_proxy():
    source_url = request.args.get('url', '').strip()
    parsed = urlparse(source_url)
    allowed_hosts = {'lh3.googleusercontent.com', 'drive.google.com', 'googleusercontent.com'}
    if parsed.scheme not in {'http', 'https'} or parsed.hostname not in allowed_hosts:
        return jsonify({'error': 'Unsupported image host'}), 400

    cached = find_cached_image(source_url)
    if cached:
        return send_file(cached, max_age=86400)
    try:
        return send_file(download_image(source_url), max_age=86400)
    except Exception as error:
        return jsonify({'error': f'Unable to load image: {error}'}), 502

@app.route('/api/upload-photo', methods=['POST'])
def upload_photo():
    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
    if file:
        filename = secure_filename(file.filename)
        file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(file_path)
        return jsonify({'photo_url': f"http://localhost:5000/uploads/{filename}"})

# --- BID HISTORY APIS ---
@app.route('/api/bids/add', methods=['POST'])
def add_bid_log():
    data = request.json
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute('''INSERT INTO bids_history (player_id, team_name, bid_amount) VALUES (?, ?, ?)''',
              (data['player_id'], data['team_name'], data['bid_amount']))
    conn.commit()
    conn.close()
    return jsonify({'message': 'Bid logged successfully'})

@app.route('/api/bids/all', methods=['GET'])
def get_all_bids():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    c.execute("SELECT * FROM bids_history ORDER BY id ASC")
    logs = [dict(row) for row in c.fetchall()]
    conn.close()
    return jsonify(logs)

@app.route('/api/bids/player/<int:player_id>', methods=['GET'])
def get_player_bids(player_id):
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    c.execute("SELECT * FROM bids_history WHERE player_id=? ORDER BY id DESC", (player_id,))
    logs = [dict(row) for row in c.fetchall()]
    conn.close()
    return jsonify(logs)

# --- PLAYERS APIS ---
@app.route('/api/players', methods=['GET'])
def get_players():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    c.execute("""
        SELECT p.*,
               COALESCE(
                   p.sold_at,
                   (SELECT MAX(b.timestamp) FROM bids_history b WHERE b.player_id = p.id)
               ) AS sold_at
        FROM players p
        ORDER BY p.id ASC
    """)
    players = [dict(row) for row in c.fetchall()]
    for player in players:
        player['reg_no'] = format_reg_no(player['reg_no'])
    conn.close()
    return jsonify(players)

@app.route('/api/players/add', methods=['POST'])
def add_player():
    data = request.json
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    try:
        photo = transform_drive_url(data.get('photo_url', ''))
        reg_no = format_reg_no(data['reg_no'])
        c.execute('''INSERT INTO players (reg_no, name, category, batting_style, bowling_style, district, block, photo_url, base_price, status)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'PENDING')''',
              (reg_no, data['name'], data['category'], data['batting_style'], 
                   data['bowling_style'], data['district'], data.get('block', '--'), photo, data.get('base_price', 400)))
        conn.commit()
        conn.close()
        return jsonify({'message': 'Player added successfully!'})
    except Exception as e:
        conn.close()
        return jsonify({'error': str(e)}), 400

@app.route('/api/players/update', methods=['POST'])
def update_player():
    data = request.json
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    photo = transform_drive_url(data.get('photo_url', ''))
    reg_no = format_reg_no(data['reg_no'])
    c.execute('''UPDATE players SET reg_no=?, name=?, category=?, batting_style=?, bowling_style=?, district=?, block=?, base_price=?, photo_url=? WHERE id=?''',
              (reg_no, data['name'], data['category'], data['batting_style'], data['bowling_style'], data['district'], data.get('block', '--'), data['base_price'], photo, data['id']))
    conn.commit()
    conn.close()
    return jsonify({'message': 'Player updated successfully!'})

@app.route('/api/players/delete/<int:player_id>', methods=['DELETE'])
def delete_player(player_id):
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute("DELETE FROM players WHERE id=?", (player_id,))
    conn.commit()
    conn.close()
    return jsonify({'message': 'Player deleted!'})

@app.route('/api/players/sell', methods=['POST'])
@app.route('/api/players/sell', methods=['POST'])
def sell_player():
    data = request.json or {}
    player_id = data.get('player_id')
    sold_to = data.get('sold_to')
    sold_price = float(data.get('sold_price') or 0)
    status = data.get('status', 'SOLD')

    if status not in ('SOLD', 'UNSOLD'):
        return jsonify({'error': 'Invalid status.'}), 400

    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()

    c.execute("SELECT status, sold_to FROM players WHERE id=?", (player_id,))
    row = c.fetchone()
    if not row:
        conn.close()
        return jsonify({'error': 'Player not found.'}), 404
    if row[0] == 'SOLD':
        conn.close()
        return jsonify({'error': f'Already SOLD to {row[1]}. This cannot be changed.'}), 409
    if row[0] == 'UNSOLD':
        conn.close()
        return jsonify({'error': 'Already marked UNSOLD. Wait for the next round.'}), 409

    if status == 'SOLD':
        c.execute("SELECT purse_left FROM teams WHERE name=?", (sold_to,))
        team = c.fetchone()
        if not team:
            conn.close()
            return jsonify({'error': 'Team not found.'}), 400
        if team[0] < sold_price:
            conn.close()
            return jsonify({'error': f'{sold_to} does not have enough purse.'}), 400
        c.execute("UPDATE players SET status='SOLD', sold_to=?, sold_price=?, sold_at=CURRENT_TIMESTAMP WHERE id=? AND status='PENDING'", (sold_to, sold_price, player_id))
        if c.rowcount == 0:
            conn.close()
            return jsonify({'error': 'Player was already finalized.'}), 409
        c.execute("UPDATE teams SET purse_left = purse_left - ?, squad_count = squad_count + 1 WHERE name=?", (sold_price, sold_to))
    else:
        c.execute("UPDATE players SET status='UNSOLD' WHERE id=? AND status='PENDING'", (player_id,))
        if c.rowcount == 0:
            conn.close()
            return jsonify({'error': 'Player was already finalized.'}), 409

    conn.commit()
    conn.close()
    return jsonify({'message': f'Player status updated to {status} successfully!'})

# --- TEAMS APIS ---
@app.route('/api/teams', methods=['GET'])
def get_teams():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    c.execute("SELECT * FROM teams")
    teams = [dict(row) for row in c.fetchall()]
    conn.close()
    return jsonify(teams)

@app.route('/api/teams/add', methods=['POST'])
def add_team():
    data = request.json
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    try:
        budget = float(data.get('budget', 50000))
        c.execute('''INSERT INTO teams (name, budget, purse_left) VALUES (?, ?, ?)''', (data['name'], budget, budget))
        conn.commit()
        conn.close()
        return jsonify({'message': 'Team added successfully!'})
    except Exception as e:
        conn.close()
        return jsonify({'error': 'Team name already exists'}), 400

@app.route('/api/teams/edit', methods=['POST'])
def edit_team():
    data = request.json
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute('''UPDATE teams SET name=?, budget=?, purse_left=? WHERE id=?''', 
              (data['name'], data['budget'], data['purse_left'], data['id']))
    conn.commit()
    conn.close()
    return jsonify({'message': 'Team updated!'})

@app.route('/api/teams/delete/<int:team_id>', methods=['DELETE'])
def delete_team(team_id):
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute("DELETE FROM teams WHERE id=?", (team_id,))
    conn.commit()
    conn.close()
    return jsonify({'message': 'Team deleted successfully!'})

# --- SETTINGS APIS ---
@app.route('/api/settings', methods=['GET', 'POST'])
def handle_settings():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    if request.method == 'POST':
        data = request.json
        c.execute('''UPDATE auction_settings SET bid_step_1=?, bid_step_2=?, bid_step_3=? WHERE id=1''',
                  (data['step1'], data['step2'], data['step3']))
        conn.commit()
        conn.close()
        return jsonify({'message': 'Bid increments updated successfully!'})
    
    c.execute("SELECT * FROM auction_settings WHERE id=1")
    settings = dict(c.fetchone())
    conn.close()
    return jsonify(settings)

# --- EXCEL IMPORT (UPDATED WITH REG_NO, BLOCK, DISTRICT, BOWLING STYLE) ---
@app.route('/api/admin/import-excel', methods=['POST'])
def import_excel():
    file = request.files.get('file')
    clear_existing = request.form.get('clear_existing') == 'true'
    
    if not file:
        return jsonify({'error': 'No file uploaded'}), 400
    
    excel_file = pd.ExcelFile(file)
    df = pd.read_excel(excel_file, sheet_name=excel_file.sheet_names[0])
    
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()

    if clear_existing:
        c.execute("DELETE FROM players")

    imported_count = 0
    used_reg_numbers = set() if clear_existing else {
        format_reg_no(row[0]) for row in c.execute("SELECT reg_no FROM players WHERE reg_no IS NOT NULL AND TRIM(reg_no) != ''")
    }
    for idx, row in df.iterrows():
        raw_reg_no = get_col_val(row, ['reg_no', 'registration_no', 'regno', 'reg_number'], '')
        # reg_no = format_reg_no(raw_reg_no) if raw_reg_no else next_registration_number(used_reg_numbers)
        # used_reg_numbers.add(reg_no)
        name = get_col_val(row, ['name', 'player_name'], 'Unknown')
        category = get_col_val(row, ['category', 'player_category', 'role'], 'All Rounder')
        batting = get_col_val(row, ['batting_style', 'batting', 'batting_type'], 'N/A')
        bowling = get_col_val(row, ['bowling_style', 'bowling', 'bowling_type'], 'N/A')
        district = get_col_val(row, ['district', 'dist'], 'N/A')
        block = get_col_val(row, ['block', 'block_name', 'salboni block player'], 'N/A')
        
        photo_raw = get_col_val(row, ['photo_url', 'player_photo', 'photo', 'image'], '')
        photo = transform_drive_url(photo_raw)
        existing = c.execute("SELECT reg_no FROM players WHERE photo_url=?", (photo,)).fetchone() if (photo and not raw_reg_no) else None
        if raw_reg_no:
            reg_no = format_reg_no(raw_reg_no)
        elif existing:
            reg_no = existing[0]
        else:
            reg_no = next_registration_number(used_reg_numbers)
        used_reg_numbers.add(reg_no)
        
        base_price_val = get_col_val(row, ['base_price', 'price', 'baseprice'], '400')
        try:
            base_price = float(base_price_val)
        except ValueError:
            base_price = 400.0

        c.execute('''INSERT INTO players (reg_no, name, category, batting_style, bowling_style, district, block, photo_url, base_price, status)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'PENDING')
                    ON CONFLICT(reg_no) DO UPDATE SET
                    name=excluded.name, category=excluded.category, batting_style=excluded.batting_style,
                    bowling_style=excluded.bowling_style, district=excluded.district, block=excluded.block,
                    photo_url=excluded.photo_url, base_price=excluded.base_price''',
                  (reg_no, name, category, batting, bowling, district, block, photo, base_price))
        imported_count += 1

    conn.commit()
    conn.close()
    return jsonify({'message': f'Successfully imported {imported_count} players!'})

    
# --- EXPORT COMPLETE AUCTION REPORT TO EXCEL ---
def style_report_sheet(ws, money_headers=(), filter_on=True):
    header_fill = PatternFill('solid', fgColor='1F2937')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(vertical='center')
    ws.freeze_panes = 'A2'
    if filter_on and ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        header = column[0].value
        longest = max((len(str(cell.value)) for cell in column if cell.value is not None), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(max(longest + 2, 10), 45)
        if header in money_headers:
            for cell in column[1:]:
                cell.number_format = '"₹"#,##0'

@app.route('/api/export/summary', methods=['GET'])
def export_summary():
    conn = sqlite3.connect(DB_NAME)
    players = pd.read_sql_query("SELECT * FROM players", conn)
    teams = pd.read_sql_query("SELECT * FROM teams", conn)
    bids = pd.read_sql_query("SELECT * FROM bids_history", conn)
    row = conn.execute("SELECT current_round FROM auction_settings WHERE id=1").fetchone()
    conn.close()
    current_round = row[0] if row else 1

    players['reg_no'] = players['reg_no'].map(format_reg_no)
    players = players.sort_values('reg_no').reset_index(drop=True)
    sold = players[players['status'] == 'SOLD']
    bid_count = bids.groupby('player_id').size()

    # Players sheet
    players_sheet = pd.DataFrame({
        'Reg No': players['reg_no'], 'Name': players['name'], 'Category': players['category'],
        'Batting': players['batting_style'], 'Bowling': players['bowling_style'],
        'District': players['district'], 'Block(Salboni)': players['block'],
        'Base Price': players['base_price'], 'Status': players['status'],
        'Sold To': players['sold_to'],
        'Sold Price': players['sold_price'].where(players['status'] == 'SOLD'),
        'Total Bids': players['id'].map(bid_count).fillna(0).astype(int),
        'Round': players['round_no'], 'Photo Link': players['photo_url'],
    })

    # Teams sheet
    spent = sold.groupby('sold_to')['sold_price'].sum()
    bought = sold.groupby('sold_to').size()
    costliest = sold.sort_values('sold_price', ascending=False).drop_duplicates('sold_to').set_index('sold_to')
    teams_sheet = pd.DataFrame({
        'Team': teams['name'], 'Budget': teams['budget'],
        'Spent': teams['name'].map(spent).fillna(0),
        'Purse Left': teams['purse_left'],
        'Players Bought': teams['name'].map(bought).fillna(0).astype(int),
    })
    teams_sheet['Average Price'] = (teams_sheet['Spent'] / teams_sheet['Players Bought'].where(teams_sheet['Players Bought'] > 0)).round(0).fillna(0)
    teams_sheet['Costliest Player'] = teams['name'].map(costliest['name']).fillna('-')
    teams_sheet['Costliest Price'] = teams['name'].map(costliest['sold_price']).fillna(0)

    # Team squads sheet (who bought whom)
    squads = sold.sort_values(['sold_to', 'sold_price'], ascending=[True, False])
    squads_sheet = pd.DataFrame({
        'Team': squads['sold_to'], 'Reg No': squads['reg_no'], 'Player': squads['name'],
        'Category': squads['category'], 'Sold Price': squads['sold_price'],
        'Base Price': squads['base_price'], 'Round': squads['round_no'],
    })

    # Bid history sheet (every bid, grouped player by player)
    lookup = players.set_index('id')
    history = bids.sort_values('id').copy()
    history['Reg No'] = history['player_id'].map(lookup['reg_no'])
    history['Player'] = history['player_id'].map(lookup['name'])
    history['Bid #'] = history.groupby('player_id').cumcount() + 1
    is_last = history['id'] == history.groupby('player_id')['id'].transform('max')
    is_sold = history['player_id'].map(lookup['status']) == 'SOLD'
    history['Result'] = ''
    history.loc[is_last & is_sold, 'Result'] = 'SOLD (winning bid)'
    history['Time (IST)'] = pd.to_datetime(history['timestamp']) + pd.Timedelta(hours=5, minutes=30)
    history = history.sort_values(['Reg No', 'id'])
    bids_sheet = history.rename(columns={'team_name': 'Team', 'bid_amount': 'Bid Amount'})[
        ['Reg No', 'Player', 'Bid #', 'Team', 'Bid Amount', 'Result', 'Time (IST)']]

    # Summary sheet
    top = sold.sort_values('sold_price', ascending=False).head(1)
    summary_sheet = pd.DataFrame([
        ('Report generated', datetime.now().strftime('%d %b %Y, %I:%M %p')),
        ('Current round', current_round),
        ('Total players', len(players)),
        ('Sold', len(sold)),
        ('Unsold (waiting for next round)', int((players['status'] == 'UNSOLD').sum())),
        ('Not yet auctioned', int((players['status'] == 'PENDING').sum())),
        ('Total money spent', float(sold['sold_price'].sum())),
        ('Average sold price', round(float(sold['sold_price'].mean()), 0) if len(sold) else 0),
        ('Costliest player', f"{top.iloc[0]['name']} - {int(top.iloc[0]['sold_price'])} ({top.iloc[0]['sold_to']})" if len(top) else '-'),
        ('Total bids placed', len(bids)),
    ], columns=['Item', 'Value'])

    output = io.BytesIO()
    money = {'Budget', 'Spent', 'Purse Left', 'Average Price', 'Costliest Price', 'Sold Price', 'Base Price', 'Bid Amount'}
    with pd.ExcelWriter(output, engine='openpyxl', datetime_format='dd-mmm-yyyy hh:mm:ss') as writer:
        summary_sheet.to_excel(writer, sheet_name='Summary', index=False)
        teams_sheet.to_excel(writer, sheet_name='Teams', index=False)
        squads_sheet.to_excel(writer, sheet_name='Team Squads', index=False)
        players_sheet.to_excel(writer, sheet_name='Players', index=False)
        bids_sheet.to_excel(writer, sheet_name='Bid History', index=False)
        for name, sheet in writer.sheets.items():
            style_report_sheet(sheet, money_headers=money, filter_on=(name != 'Summary'))

    output.seek(0)
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'Auction_Report_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx'
    )

@app.route('/api/admin/start-next-round', methods=['POST'])
def next_round():
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute("UPDATE players SET status='PENDING', round_no = round_no + 1 WHERE status='UNSOLD'")
    moved = c.rowcount
    if moved == 0:
        conn.close()
        return jsonify({'message': 'No unsold players to move to the next round.'})
    c.execute("UPDATE auction_settings SET current_round = current_round + 1 WHERE id=1")
    c.execute("SELECT current_round FROM auction_settings WHERE id=1")
    round_no = c.fetchone()[0]
    conn.commit()
    conn.close()
    return jsonify({'message': f'Round {round_no} started. {moved} unsold players are back in the auction pool.'})

@app.route('/api/admin/cache-images', methods=['POST'])
def cache_all_images():
    conn = sqlite3.connect(DB_NAME)
    rows = conn.execute("SELECT name, photo_url FROM players").fetchall()
    conn.close()

    already, no_link, todo = 0, 0, {}
    for name, photo_url in rows:
        source = drive_source_url(photo_url)
        if not source:
            no_link += 1
        elif find_cached_image(source):
            already += 1
        else:
            todo.setdefault(source, name)   # same image used twice is downloaded only once

    def work(item):
        source, name = item
        try:
            download_image(source)
            return None
        except Exception as error:
            return f'{name}: {error}'

    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(work, todo.items()))
    failed = [r for r in results if r]
    downloaded = len(results) - len(failed)

    message = f'Photos saved. New: {downloaded}, already saved: {already}, failed: {len(failed)}, no Drive link: {no_link}.'
    if failed:
        message += ' First failure - ' + failed[0]
    return jsonify({'message': message})

@app.route('/api/admin/reset-auction', methods=['POST'])
def reset_auction():
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute("UPDATE players SET status='PENDING', sold_to=NULL, sold_price=0, round_no=1")
    c.execute("UPDATE teams SET purse_left=budget, squad_count=0")
    c.execute("DELETE FROM bids_history")
    c.execute("UPDATE auction_settings SET current_round=1 WHERE id=1")
    conn.commit()
    conn.close()
    return jsonify({'message': 'Auction reset. All players are PENDING again and all team purses are restored.'})

if __name__ == '__main__':
    app.run(port=5000, debug=True)